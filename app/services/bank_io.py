"""题库 Excel 导入导出（导入/导出格式一致）。"""

import io
import json
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.ai_generator import Q_TYPE_LABELS

META_SHEET = "题库"
QUESTIONS_SHEET = "题目"

META_FIELDS = ("题库名称", "分类", "说明")
QUESTION_HEADERS = (
    "序号",
    "题型",
    "题干",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "选项F",
    "答案",
    "解析",
)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_TYPE_TO_LABEL = {k: v for k, v in Q_TYPE_LABELS.items()}
_LABEL_TO_TYPE = {v: k for k, v in Q_TYPE_LABELS.items()}
_LABEL_TO_TYPE.update(
    {
        "单选": "single",
        "多选": "multiple",
        "判断": "judge",
        "single": "single",
        "multiple": "multiple",
        "judge": "judge",
    }
)


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _normalize_answer(ans: str) -> str:
    raw = (ans or "").upper().replace(" ", "")
    if not raw:
        return ""
    if "," in raw:
        return ",".join(sorted(p for p in raw.split(",") if p))
    return "".join(sorted(raw))


def _parse_q_type(raw: str) -> str:
    s = _cell_str(raw)
    if not s:
        return "single"
    if s in _LABEL_TO_TYPE:
        return _LABEL_TO_TYPE[s]
    lower = s.lower()
    if lower in _LABEL_TO_TYPE:
        return _LABEL_TO_TYPE[lower]
    return "single"


def _options_from_row(row: tuple) -> list[str]:
    """row: 0-based, indices 3-8 are 选项A-F (1-based col 4-9)."""
    opts = []
    for i in range(3, 9):
        if i < len(row):
            text = _cell_str(row[i])
            if text:
                opts.append(text)
    return opts


def build_bank_payload(
    title: str,
    category: str,
    description: str,
    questions: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "title": title,
        "category": category,
        "description": description,
        "questions": questions,
    }


def _style_header_row(ws, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def export_bank_xlsx(
    title: str,
    category: str,
    description: str,
    questions: list[Any],
) -> bytes:
    wb = Workbook()
    meta = wb.active
    meta.title = META_SHEET
    meta["A1"] = "字段"
    meta["B1"] = "值"
    _style_header_row(meta, 2)
    meta.column_dimensions["B"].width = 50
    meta_values = {
        "题库名称": title or "未命名题库",
        "分类": category or "未分类",
        "说明": description or "",
    }
    for i, field in enumerate(META_FIELDS, start=2):
        meta[f"A{i}"] = field
        meta[f"B{i}"] = meta_values.get(field, "")

    ws = wb.create_sheet(QUESTIONS_SHEET)
    for col, header in enumerate(QUESTION_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header)
    _style_header_row(ws, len(QUESTION_HEADERS))
    ws.column_dimensions[get_column_letter(3)].width = 50
    ws.column_dimensions[get_column_letter(11)].width = 50

    for idx, q in enumerate(questions, start=1):
        q_type = getattr(q, "q_type", None) or q.get("q_type", "single")
        content = getattr(q, "content", None) or q.get("content", "")
        options = getattr(q, "options", None) or q.get("options", []) or []
        answer = getattr(q, "answer", None) or q.get("answer", "")
        explanation = getattr(q, "explanation", None) or q.get("explanation", "")

        if not isinstance(options, list):
            options = []

        row_data = [
            idx,
            _TYPE_TO_LABEL.get(q_type, q_type),
            content,
            *[(options[i] if i < len(options) else "") for i in range(6)],
            answer,
            explanation,
        ]
        for col, val in enumerate(row_data, start=1):
            ws.cell(row=idx + 1, column=col, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_meta_sheet(ws) -> dict[str, str]:
    meta = {"title": "导入题库", "category": "未分类", "description": ""}
    field_map = {
        "题库名称": "title",
        "分类": "category",
        "说明": "description",
    }
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = _cell_str(row[0])
        val = _cell_str(row[1]) if len(row) > 1 else ""
        if key in field_map and val:
            meta[field_map[key]] = val
    return meta


def _read_questions_sheet(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [_cell_str(c) for c in rows[0]]
    start = 1
    if header and header[0] in ("序号", "No", "no", "#"):
        start = 1
    else:
        start = 0

    questions = []
    for row in rows[start:]:
        if not row or not any(row):
            continue
        content = _cell_str(row[2]) if len(row) > 2 else ""
        if not content:
            continue
        q_type = _parse_q_type(row[1] if len(row) > 1 else "")
        options = _options_from_row(row)
        if q_type == "judge" and len(options) < 2:
            options = ["正确", "错误"]
        if q_type != "judge" and len(options) < 2:
            continue
        answer = _cell_str(row[9]) if len(row) > 9 else ""
        if not answer:
            continue
        explanation = _cell_str(row[10]) if len(row) > 10 else ""
        questions.append(
            {
                "q_type": q_type,
                "content": content,
                "options": options,
                "answer": _normalize_answer(answer),
                "explanation": explanation,
            }
        )
    return questions


def import_bank_xlsx(data: bytes) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if META_SHEET in wb.sheetnames:
        meta = _read_meta_sheet(wb[META_SHEET])
    else:
        meta = {"title": "导入题库", "category": "未分类", "description": ""}

    if QUESTIONS_SHEET in wb.sheetnames:
        questions = _read_questions_sheet(wb[QUESTIONS_SHEET])
    else:
        questions = _read_questions_sheet(wb.active)

    wb.close()
    if not questions:
        raise ValueError("未找到有效题目，请检查「题目」工作表格式")
    return build_bank_payload(
        meta["title"],
        meta["category"],
        meta["description"],
        questions,
    )


def import_bank_json(data: bytes) -> dict[str, Any]:
    """兼容旧版 JSON 导出文件。"""
    obj = json.loads(data.decode("utf-8"))
    title = obj.get("title") or "导入题库"
    category = obj.get("category") or "未分类"
    description = obj.get("description") or ""
    raw_qs = obj.get("questions") or []
    questions = []
    for item in raw_qs:
        if not isinstance(item, dict):
            continue
        content = _cell_str(item.get("content"))
        if not content:
            continue
        q_type = _parse_q_type(item.get("q_type", "single"))
        options = item.get("options") or []
        if not isinstance(options, list):
            options = []
        options = [_cell_str(o) for o in options if _cell_str(o)]
        if q_type == "judge" and len(options) < 2:
            options = ["正确", "错误"]
        if len(options) < 2:
            continue
        answer = _cell_str(item.get("answer"))
        if not answer:
            continue
        questions.append(
            {
                "q_type": q_type,
                "content": content,
                "options": options,
                "answer": _normalize_answer(answer),
                "explanation": _cell_str(item.get("explanation")),
            }
        )
    if not questions:
        raise ValueError("JSON 中无有效题目")
    return build_bank_payload(title, category, description, questions)


def parse_import_file(filename: str, data: bytes) -> dict[str, Any]:
    lower = (filename or "").lower()
    if lower.endswith(".json"):
        return import_bank_json(data)
    if lower.endswith((".xlsx", ".xlsm")):
        return import_bank_xlsx(data)
    raise ValueError("仅支持 .xlsx 或 .json 文件（推荐使用 Excel 格式）")
