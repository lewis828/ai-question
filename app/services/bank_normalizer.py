"""将非标准格式的题库文档解析并规范化为标准 Excel 结构。"""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.ai_generator import finalize_judge_question, normalize_judge_answer
from app.services.bank_io import build_bank_payload, import_bank_xlsx
from app.services.parser import extract_text

TYPE_ORDER = ("single", "judge", "multiple")
_LABELS = ("A", "B", "C", "D", "E", "F")

_TYPE_MAP = {
    "单选题": "single",
    "单选": "single",
    "多选题": "multiple",
    "多选": "multiple",
    "判断题": "judge",
    "判断": "judge",
    "single": "single",
    "multiple": "multiple",
    "judge": "judge",
}

_QUESTION_START = re.compile(
    r"^(?:【?(?P<type>单选题|多选题|判断题|单选|多选|判断)【?】?\s*)?"
    r"(?:(?:第\s*)?(?P<num>\d+)\s*(?:题)?[\.\．、:：\)）\]]\s*)?"
    r"(?P<stem>.+)$"
)
_OPTION = re.compile(r"^([A-Fa-f])[\.．、:：\)]\s*(.+)$")
_ANSWER = re.compile(r"^(?:【|\[)?(?:参考)?答案(?:】|\])?[：:]\s*(?P<ans>.+)$")
_EXPLAIN = re.compile(r"^(?:【|\[)?(?:解析|说明)(?:】|\])?[：:]\s*(?P<exp>.+)$")
_TYPE_TAG = re.compile(r"【?(单选题|多选题|判断题|单选|多选|判断)】?")
_NUM_START = re.compile(r"^(?:【?(?:单选题|多选题|判断题|单选|多选|判断)】?\s*)?(?:第\s*)?(\d+)\s*(?:题)?[\.\．、:：\)）\]]\s*\S")


def _is_question_start(line: str) -> bool:
    if _OPTION.match(line):
        return False
    if _ANSWER.match(line) or line.startswith(("答案", "【答案", "[答案", "参考答案")):
        return False
    if _EXPLAIN.match(line) or line.startswith(("解析", "【解析", "[解析", "说明", "【说明")):
        return False
    if _NUM_START.match(line):
        return True
    if re.match(r"^【(?:单选题|多选题|判断题|单选|多选|判断)】", line):
        return True
    return False


def _clean_stem(text: str) -> str:
    s = (text or "").strip()
    s = _TYPE_TAG.sub("", s)
    s = re.sub(r"^第\s*\d+\s*题[：:\.\s]*", "", s)
    s = re.sub(r"^\d+[\.\．、\)）\]]\s*", "", s)
    s = re.sub(r"^(?:判断题|单选题|多选题|单选|多选|判断)[：:\s]*", "", s)
    return s.strip()


def _normalize_judge_answer(raw: str, options: list[str] | None = None) -> str:
    return normalize_judge_answer(raw)


def _normalize_choice_answer(raw: str) -> str:
    letters = re.findall(r"[A-Fa-f]", raw.upper())
    if not letters:
        return ""
    if len(letters) == 1:
        return letters[0]
    return ",".join(sorted(set(letters)))


def _infer_q_type(explicit: str, options: list[str], answer: str) -> str:
    if explicit in _TYPE_MAP:
        return _TYPE_MAP[explicit]
    if options == ["正确", "错误"] or set(options) >= {"正确", "错误"}:
        return "judge"
    ans = answer.replace(",", "")
    if len(ans) > 1:
        return "multiple"
    return "single"


def _split_blocks(lines: list[str]) -> list[list[str]]:
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if _is_question_start(line):
            if current:
                blocks.append(current)
            current = [line]
        else:
            current.append(line)
    if current:
        blocks.append(current)
    return blocks


def _parse_block(block: list[str]) -> dict[str, Any] | None:
    if not block:
        return None

    first = block[0]
    m = _QUESTION_START.match(first)
    explicit_type = m.group("type") if m else ""
    if not explicit_type:
        tag = _TYPE_TAG.search(first)
        explicit_type = tag.group(1) if tag else ""
    stem = _clean_stem(m.group("stem") if m else first)

    options: list[str] = []
    answer = ""
    explanation = ""
    extra_stem: list[str] = []

    for line in block[1:]:
        opt = _OPTION.match(line)
        if opt:
            options.append(opt.group(2).strip())
            continue

        ans = _ANSWER.match(line)
        if ans:
            answer = ans.group("ans").strip()
            continue

        exp = _EXPLAIN.match(line)
        if exp:
            explanation = exp.group("exp").strip()
            continue

        if line.startswith(("解析", "说明", "【解析", "[解析")):
            explanation = re.sub(r"^(?:【|\[)?(?:解析|说明)(?:】|\])?[：:\s]*", "", line).strip()
            continue

        if not options and not answer:
            extra_stem.append(line)

    if extra_stem:
        stem = _clean_stem(" ".join([stem, *extra_stem]))

    if not stem:
        return None

    if not options:
        if explicit_type in {"判断", "判断题"} or "判断" in (explicit_type or ""):
            options = ["正确", "错误"]
        else:
            return None

    q_type = _infer_q_type(explicit_type or "", options, answer)
    if q_type == "judge":
        answer = _normalize_judge_answer(answer or "正确", options)
        options = ["正确", "错误"]
    else:
        answer = _normalize_choice_answer(answer)
        if not answer:
            return None

    if not explanation:
        explanation = _default_explanation(q_type, answer, options)

    return {
        "q_type": q_type,
        "content": stem,
        "options": options[:6],
        "answer": answer,
        "explanation": explanation,
    }


def _default_explanation(q_type: str, answer: str, options: list[str]) -> str:
    if q_type == "judge":
        from app.services.ai_generator import build_judge_explanation

        return build_judge_explanation(answer)
    labels = [p for p in re.split(r"[,，]", answer) if p in _LABELS]
    correct = "、".join(options[_LABELS.index(l)] for l in labels if l in _LABELS[: len(options)])
    return f"正确答案 {answer.replace(',', '、')}：{correct}。"


def parse_questions_from_text(text: str) -> list[dict[str, Any]]:
    lines = [ln.strip() for ln in text.replace("\r\n", "\n").split("\n")]
    lines = [ln for ln in lines if ln]
    if not lines:
        return []

    if not any(_is_question_start(ln) for ln in lines):
        joined = "\n".join(lines)
        parts = re.split(r"(?=(?:^|\n)\s*(?:\d+[\.\、]|【?(?:单选|多选|判断)))", joined)
        lines = [p.strip() for p in parts if p.strip()]

    blocks = _split_blocks(lines)
    questions: list[dict[str, Any]] = []
    for block in blocks:
        q = _parse_block(block)
        if q:
            questions.append(q)
    return _sort_questions(questions)


def _sort_questions(questions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    order = {t: i for i, t in enumerate(TYPE_ORDER)}
    sorted_qs = sorted(questions, key=lambda q: order.get(q.get("q_type", "single"), 99))
    return [
        finalize_judge_question(q) if q.get("q_type") == "judge" else q
        for q in sorted_qs
    ]


def _parse_flexible_xlsx(path: Path) -> dict[str, Any]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("Excel 为空")

    header = [str(c or "").strip() for c in rows[0]]

    def col(*names: str) -> int | None:
        for i, h in enumerate(header):
            for name in names:
                if name in h:
                    return i
        return None

    idx_type = col("题型")
    idx_stem = col("题干", "题目", "问题")
    idx_answer = col("答案")
    idx_explain = col("解析", "说明")
    option_cols = [i for i, h in enumerate(header) if h.startswith("选项") or re.match(r"^[A-F]$", h)]

    if idx_stem is None:
        raise ValueError("无法识别 Excel 列结构，请使用规范格式或 Word/TXT 题库")

    title = path.stem[:200] or "规范题库"
    questions: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not row or not row[idx_stem]:
            continue
        stem = _clean_stem(str(row[idx_stem]))
        options = []
        if option_cols:
            for i in option_cols:
                if i < len(row) and row[i]:
                    options.append(str(row[i]).strip())
        else:
            for i in range(idx_stem + 1, min(idx_stem + 7, len(row))):
                if i == idx_answer or i == idx_explain:
                    continue
                if row[i]:
                    options.append(str(row[i]).strip())

        answer_raw = str(row[idx_answer]).strip() if idx_answer is not None and idx_answer < len(row) and row[idx_answer] else ""
        explicit = str(row[idx_type]).strip() if idx_type is not None and idx_type < len(row) and row[idx_type] else ""
        q_type = _infer_q_type(explicit, options, answer_raw)
        if q_type == "judge":
            options = ["正确", "错误"]
            answer = _normalize_judge_answer(answer_raw, options)
        else:
            answer = _normalize_choice_answer(answer_raw)
        if not answer:
            continue

        explanation = ""
        if idx_explain is not None and idx_explain < len(row) and row[idx_explain]:
            explanation = str(row[idx_explain]).strip()
        if not explanation:
            explanation = _default_explanation(q_type, answer, options)

        questions.append(
            {
                "q_type": q_type,
                "content": stem,
                "options": options[:6],
                "answer": answer,
                "explanation": explanation,
            }
        )

    if not questions:
        raise ValueError("未能从 Excel 中解析有效题目")
    return build_bank_payload(title, "未分类", f"规范转换 · {len(questions)} 道题", _sort_questions(questions))


def normalize_bank_from_file(
    path: Path,
    *,
    title: str = "",
    category: str = "未分类",
) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            payload = import_bank_xlsx(path.read_bytes())
        except ValueError:
            payload = _parse_flexible_xlsx(path)
        if title:
            payload["title"] = title
        if category and category != "未分类":
            payload["category"] = category
        payload["questions"] = _sort_questions(payload.get("questions") or [])
        return payload

    if suffix not in {".pdf", ".doc", ".docx", ".txt", ".md"}:
        raise ValueError(f"不支持格式: {suffix}")

    text = extract_text(path)
    questions = parse_questions_from_text(text)
    if not questions:
        raise ValueError("未能从资料中识别题库结构，请检查题目、选项、答案格式")

    bank_title = title or path.stem[:200] or "规范题库"
    return build_bank_payload(
        bank_title,
        category,
        f"规范转换 · {len(questions)} 道题",
        questions,
    )


def normalize_bank_from_bytes(filename: str, data: bytes, **kwargs: Any) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower() or ".txt"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        tmp_path = Path(tmp.name)
    try:
        return normalize_bank_from_file(tmp_path, **kwargs)
    finally:
        tmp_path.unlink(missing_ok=True)
