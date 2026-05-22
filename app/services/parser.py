from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".md", ".pptx", ".ppt", ".xlsx", ".xls"}


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"不支持的文件格式: {suffix}")

    if suffix == ".pdf":
        return _parse_pdf(file_path)
    if suffix in {".docx", ".doc"}:
        return _parse_docx(file_path)
    if suffix in {".pptx", ".ppt"}:
        return _parse_pptx(file_path)
    if suffix in {".xlsx", ".xls"}:
        return _parse_xlsx(file_path)
    return file_path.read_text(encoding="utf-8", errors="ignore")


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text.strip())
    return "\n\n".join(parts)


def _parse_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())


def _parse_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return "\n".join(parts)


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)
